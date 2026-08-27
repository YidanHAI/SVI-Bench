#!/usr/bin/env python3
"""Read Judge workbook task specs and build the versioned rubric contract."""

import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


RUBRIC_VERSION = "workbook-v2"
DEFAULT_TASK_SHEET = "V1_题目池"
DEFAULT_RUBRIC_SHEET = "评测维度"
DIMENSION_IDS = ("D1", "D2", "D3", "D4", "D5")
REQUIRED_RUBRIC_HEADERS = ("维度", "定义", "触发条件", "与产品能力的对应")

_DATA_VALIDATION_ID = re.compile(
    rb'(<(?:[A-Za-z_][\w.-]*:)?dataValidation\b[^>]*?)\s+id="[^"]*"'
)
_DIMENSION_LABEL = re.compile(r"^(D[1-5])\s*[.．、]?\s*(.*)$", re.IGNORECASE)


def clean(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, str):
        value = value.replace("\r\n", "\n").replace("\r", "\n").strip()
        return value or None
    return value


def parse_bool(value):
    return str(clean(value) or "").strip() == "是"


def parse_weight(value):
    value = clean(value)
    if value in (None, "-", "—"):
        return None
    return float(value)


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(8 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def canonical_sha256(value):
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _strip_wps_data_validation_ids(raw):
    """Return an in-memory XLSX without WPS-only dataValidation ``id`` attrs."""
    source_buffer = io.BytesIO(raw)
    target_buffer = io.BytesIO()
    removed = 0
    with zipfile.ZipFile(source_buffer, "r") as source:
        with zipfile.ZipFile(target_buffer, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/") and payload:
                    payload, count = _DATA_VALIDATION_ID.subn(rb"\1", payload)
                    removed += count
                target.writestr(info, payload)
    return target_buffer.getvalue(), removed


def load_workbook_compatible(path, **kwargs):
    """Load Excel while tolerating WPS's non-standard dataValidation ``id``."""
    path = Path(path)
    raw = path.read_bytes()
    sanitized, removed = _strip_wps_data_validation_ids(raw)
    if removed:
        buffer = io.BytesIO(sanitized)
        workbook = load_workbook(buffer, **kwargs)
        # Keep the buffer alive for read-only workbooks, whose ZipFile is lazy.
        workbook._joyvl_wps_compat_buffer = buffer
        workbook._joyvl_wps_data_validation_ids_removed = removed
        return workbook
    workbook = load_workbook(path, **kwargs)
    workbook._joyvl_wps_data_validation_ids_removed = 0
    return workbook


def _sheet_rows(sheet):
    for row_index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
        yield row_index, [clean(value) for value in values]


def load_task_specs(path, sheet_name=DEFAULT_TASK_SHEET):
    workbook = load_workbook_compatible(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"Missing sheet {sheet_name!r} in {path}")
        sheet = workbook[sheet_name]
        rows = _sheet_rows(sheet)
        try:
            _, headers = next(rows)
        except StopIteration as exc:
            raise RuntimeError(f"Sheet {sheet_name!r} is empty in {path}") from exc
        specs = {}
        for row_index, values in rows:
            row = {
                header: clean(values[column]) if column < len(values) else None
                for column, header in enumerate(headers)
                if header
            }
            task_id = str(row.get("id") or "").strip()
            if not task_id or str(row.get("入选状态") or "").strip() != "入选":
                continue
            dimensions = {}
            for name in DIMENSION_IDS:
                dimensions[name] = {
                    "applicable": parse_bool(row.get(f"{name} 适用")),
                    "weight": parse_weight(row.get(f"{name} 权重")),
                    "threshold": clean(row.get(f"{name} G/S/B 阈值")) or "-",
                }
            if task_id in specs:
                raise RuntimeError(
                    f"Duplicate selected task id {task_id!r} in rows "
                    f"{specs[task_id]['source_row']} and {row_index}"
                )
            specs[task_id] = {
                "id": task_id,
                "category": row.get("分类"),
                "scene": row.get("场景标签"),
                "user_prompt": row.get("用户 query"),
                "trigger": row.get("触发点(文字描述)"),
                "expected_response_window": row.get("期望响应窗口"),
                "expected_key_information": row.get("期望关键信息"),
                "dimensions": dimensions,
                "source_row": row_index,
            }
        return specs
    finally:
        workbook.close()


def _issue(code, message, severity="error", **details):
    return {
        "severity": severity,
        "code": code,
        "message": message,
        **details,
    }


def load_rubric_contract(path, sheet_name=DEFAULT_RUBRIC_SHEET):
    workbook = load_workbook_compatible(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"Missing sheet {sheet_name!r} in {path}")
        sheet = workbook[sheet_name]
        rows = list(_sheet_rows(sheet))
    finally:
        workbook.close()

    if not rows:
        raise RuntimeError(f"Sheet {sheet_name!r} is empty in {path}")
    headers = rows[0][1]
    issues = []
    header_index = {header: index for index, header in enumerate(headers) if header}
    missing_headers = [name for name in REQUIRED_RUBRIC_HEADERS if name not in header_index]
    if missing_headers:
        issues.append(_issue(
            "rubric_missing_headers",
            f"Rubric sheet is missing required headers: {missing_headers}",
            sheet=sheet_name,
            headers=missing_headers,
        ))

    dimensions = {}
    for row_index, values in rows[1:]:
        for column_index, value in enumerate(values, start=1):
            header = headers[column_index - 1] if column_index <= len(headers) else None
            if value is not None and not header:
                cell = f"{get_column_letter(column_index)}{row_index}"
                issues.append(_issue(
                    "rubric_headerless_value",
                    f"Unexpected value outside the rubric table at {sheet_name}!{cell}",
                    severity="warning",
                    sheet=sheet_name,
                    cell=cell,
                    value=value,
                ))

        def field(name):
            index = header_index.get(name)
            return values[index] if index is not None and index < len(values) else None

        raw_label = field("维度")
        if raw_label is None:
            continue
        match = _DIMENSION_LABEL.match(str(raw_label).strip())
        if not match:
            issues.append(_issue(
                "rubric_invalid_dimension_label",
                f"Cannot parse dimension label {raw_label!r}",
                sheet=sheet_name,
                row=row_index,
                value=raw_label,
            ))
            continue
        dimension_id = match.group(1).upper()
        if dimension_id in dimensions:
            issues.append(_issue(
                "rubric_duplicate_dimension",
                f"Dimension {dimension_id} appears more than once",
                sheet=sheet_name,
                row=row_index,
                dimension=dimension_id,
            ))
            continue
        dimensions[dimension_id] = {
            "name": clean(match.group(2)),
            "definition": field("定义"),
            "trigger_condition": field("触发条件"),
            "product_capability": field("与产品能力的对应"),
            "source_row": row_index,
        }

    missing_dimensions = [name for name in DIMENSION_IDS if name not in dimensions]
    if missing_dimensions:
        issues.append(_issue(
            "rubric_missing_dimensions",
            f"Rubric sheet is missing dimensions: {missing_dimensions}",
            sheet=sheet_name,
            dimensions=missing_dimensions,
        ))
    ordered_dimensions = {
        name: dimensions[name] for name in DIMENSION_IDS if name in dimensions
    }
    contract = {
        "version": RUBRIC_VERSION,
        "source_sheet": sheet_name,
        "dimension_order": list(DIMENSION_IDS),
        "dimensions": ordered_dimensions,
        "judge_rules": {
            "applicability_authority": "task_spec.dimensions.<dimension>.applicable",
            "threshold_authority": "task_spec.dimensions.<dimension>.threshold",
            "per_task_thresholds_override_global_descriptions": True,
            "grade_score_mapping": {"G": 1.0, "S": 0.5, "B": 0.0},
            "D3_no_valid_response_score": 0.0,
            "D4_scope": "response_content_correctness",
        },
    }
    return contract, issues


def _first_int(text, pattern):
    match = re.search(pattern, str(text or ""), flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def validate_task_specs(specs):
    """Detect known cross-field contradictions that make formal grading unsafe."""
    issues = []
    for task_id, spec in specs.items():
        applicable_weight_sum = 0.0
        for dimension_id, dimension in (spec.get("dimensions") or {}).items():
            applicable = dimension.get("applicable") is True
            weight = dimension.get("weight")
            threshold = dimension.get("threshold")
            has_rule = weight is not None or threshold not in (None, "-", "—")
            if applicable and (weight is None or threshold in (None, "-", "—")):
                issues.append(_issue(
                    "task_applicable_dimension_missing_rule",
                    f"Task {task_id} marks {dimension_id} applicable but lacks weight/threshold",
                    task_id=task_id,
                    source_row=spec.get("source_row"),
                    dimension=dimension_id,
                    dimension_spec=dimension,
                ))
            if not applicable and has_rule:
                issues.append(_issue(
                    "task_nonapplicable_dimension_has_rule",
                    f"Task {task_id} marks {dimension_id} non-applicable but keeps weight/threshold",
                    task_id=task_id,
                    source_row=spec.get("source_row"),
                    dimension=dimension_id,
                    dimension_spec=dimension,
                ))
            if applicable and weight is not None:
                applicable_weight_sum += float(weight)
        if abs(applicable_weight_sum - 1.0) > 1e-9:
            issues.append(_issue(
                "task_applicable_weight_sum",
                f"Task {task_id} applicable weights sum to {applicable_weight_sum:g}, not 1",
                severity="warning",
                task_id=task_id,
                source_row=spec.get("source_row"),
                weight_sum=applicable_weight_sum,
            ))

        deadline = _first_int(spec.get("user_prompt"), r"限时\s*(\d+)\s*(?:秒|s)″?")
        if deadline is None:
            continue
        duration_fields = {
            "user_prompt": deadline,
            "scene": _first_int(spec.get("scene"), r"[-—]\s*(\d+)\s*(?:秒|s)"),
            "expected_response_window": _first_int(
                spec.get("expected_response_window"),
                r"R1\s*后(?:的)?\s*(\d+)\s*(?:秒|s)",
            ),
            "expected_key_information": _first_int(
                spec.get("expected_key_information"),
                r"时间到\s*/\s*(\d+)\s*(?:秒|s)\s*结束",
            ),
        }
        conflicting = {
            field: value
            for field, value in duration_fields.items()
            if value is not None and value != deadline
        }
        if conflicting:
            issues.append(_issue(
                "task_duration_conflict",
                f"Task {task_id} has a {deadline}s query deadline but conflicting duration fields",
                task_id=task_id,
                source_row=spec.get("source_row"),
                durations={
                    field: value
                    for field, value in duration_fields.items()
                    if value is not None
                },
            ))
    return issues


def build_workbook_v2_contract(path, task_specs, rubric_sheet=DEFAULT_RUBRIC_SHEET):
    contract, issues = load_rubric_contract(path, rubric_sheet)
    issues.extend(validate_task_specs(task_specs))
    return {
        "rubric_version": RUBRIC_VERSION,
        "rubric_contract": contract,
        "rubric_contract_sha256": canonical_sha256(contract),
        "rubric_validation": {
            "formal_ready": not any(
                item.get("severity") == "error" for item in issues
            ),
            "issues": issues,
        },
    }


def bind_contract_to_spec(spec, contract_bundle):
    bound = dict(spec)
    bound.update({
        "rubric_version": contract_bundle["rubric_version"],
        "rubric_contract_sha256": contract_bundle["rubric_contract_sha256"],
        "rubric_contract": contract_bundle["rubric_contract"],
    })
    return bound


def format_blocking_issues(issues):
    lines = []
    for item in issues:
        if item.get("severity") != "error":
            continue
        location = item.get("cell") or item.get("task_id") or item.get("sheet") or "workbook"
        lines.append(f"- {item['code']} [{location}]: {item['message']}")
    return "\n".join(lines)
