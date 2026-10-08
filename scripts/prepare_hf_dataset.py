#!/usr/bin/env python3
"""Build a metadata-sanitized Hugging Face dataset directory."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
VIDEO_SUFFIXES = {".mp4", ".m4v", ".mov"}
EXPECTED_TASKS = 75
EXPECTED_VIDEOS = 71
SENSITIVE_CELL_PATTERN = re.compile(
    r"(?:/jpfs/|rtsp://|https?://(?:10\.|127\.0\.0\.1|localhost)|"
    r"(?:sk|ark|pk)-[A-Za-z0-9_-]{12,})",
    re.IGNORECASE,
)
SENSITIVE_METADATA_KEYS = {
    "author",
    "comment",
    "copyright",
    "creationtime",
    "description",
    "did",
    "gps",
    "location",
    "lvmetainfo",
    "synopsis",
    "uid",
    "videoid",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--video-dir", type=Path, required=True)
    parser.add_argument(
        "--pilot-labels",
        type=Path,
        help="Optional human-rating workbook used to reproduce Judge alignment.",
    )
    parser.add_argument(
        "--media-index",
        type=Path,
        default=ROOT / "dataset" / "media_index.jsonl",
    )
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--dataset-id", default="Danmel02/SVI-bench")
    parser.add_argument("--provenance", type=Path)
    parser.add_argument("--data-license", default="other")
    parser.add_argument("--finalize", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--ffmpeg", default="ffmpeg")
    parser.add_argument("--ffprobe", default="ffprobe")
    return parser.parse_args()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        if not line.strip():
            continue
        value = json.loads(line)
        if not isinstance(value, dict):
            raise ValueError(f"{path}:{line_number} must contain a JSON object")
        rows.append(value)
    return rows


def load_media_index(path: Path, video_dir: Path) -> tuple[list[dict[str, str]], dict[str, Path]]:
    rows = read_jsonl(path)
    if len(rows) != EXPECTED_TASKS:
        raise ValueError(f"media index must contain {EXPECTED_TASKS} tasks, found {len(rows)}")

    normalized = []
    sources: dict[str, Path] = {}
    seen_ids: set[str] = set()
    for row in rows:
        task_id = str(row.get("id") or "").strip()
        video = str(row.get("video") or "").strip()
        relative = PurePosixPath(video)
        if not task_id or task_id in seen_ids:
            raise ValueError(f"media index contains an empty or duplicate task id: {task_id!r}")
        if relative.is_absolute() or len(relative.parts) != 1 or ".." in relative.parts:
            raise ValueError(f"media index video must be a basename: {video!r}")
        source = (video_dir / video).resolve()
        if source.parent != video_dir.resolve() or not source.is_file():
            raise ValueError(f"mapped video is missing or outside --video-dir: {video!r}")
        seen_ids.add(task_id)
        normalized.append({"id": task_id, "video": video})
        sources[video] = source

    if len(sources) != EXPECTED_VIDEOS:
        raise ValueError(
            f"media index must reference {EXPECTED_VIDEOS} unique videos, found {len(sources)}"
        )
    available = {
        path.name
        for path in video_dir.iterdir()
        if path.is_file() and path.suffix.lower() in VIDEO_SUFFIXES and not path.name.startswith(".")
    }
    if available != set(sources):
        missing = sorted(set(sources) - available)
        unreferenced = sorted(available - set(sources))
        raise ValueError(
            f"video directory and media index differ; missing={missing}, unreferenced={unreferenced}"
        )
    return normalized, sources


def workbook_task_ids(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "题目池" not in workbook.sheetnames:
            raise ValueError("workbook is missing the 题目池 sheet")
        sheet = workbook["题目池"]
        headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
        try:
            id_column = headers.index("id") + 1
        except ValueError as exc:
            raise ValueError("题目池 is missing the id column") from exc
        ids = [
            str(sheet.cell(row, id_column).value or "").strip()
            for row in range(2, sheet.max_row + 1)
        ]
        ids = [value for value in ids if value]
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and SENSITIVE_CELL_PATTERN.search(cell.value):
                        raise ValueError(
                            f"workbook contains a private path, endpoint, or credential-like value "
                            f"at {sheet.title}!{cell.coordinate}"
                        )
        return ids
    finally:
        workbook.close()


def sanitize_workbook(source: Path, destination: Path) -> None:
    workbook = load_workbook(source, keep_links=False)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and SENSITIVE_CELL_PATTERN.search(cell.value):
                    workbook.close()
                    raise ValueError(
                        f"workbook contains a private path, endpoint, or credential-like value "
                        f"at {sheet.title}!{cell.coordinate}"
                    )
                # Comments retain author metadata independently of cell values.
                if cell.comment is not None:
                    cell.comment = None
    properties = workbook.properties
    properties.creator = "SVI-Bench"
    properties.lastModifiedBy = "SVI-Bench"
    properties.title = "SVI-Bench"
    properties.subject = "Streaming video interaction benchmark"
    properties.description = None
    properties.keywords = None
    properties.category = None
    properties.contentStatus = None
    properties.identifier = None
    properties.language = None
    properties.created = datetime(2026, 1, 1)
    properties.modified = datetime(2026, 1, 1)
    workbook.custom_doc_props.props.clear()
    workbook.save(destination)
    workbook.close()


def probe_media(path: Path, ffprobe: str) -> dict[str, Any]:
    completed = subprocess.run(
        [
            ffprobe,
            "-v",
            "error",
            "-show_entries",
            "format=duration:format_tags:stream=index,codec_type,codec_name,width,height,sample_rate,channels:stream_tags",
            "-of",
            "json",
            str(path),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    return json.loads(completed.stdout)


def media_signature(probe: dict[str, Any]) -> list[tuple[Any, ...]]:
    return sorted([
        (
            stream.get("codec_type"),
            stream.get("codec_name"),
            stream.get("width"),
            stream.get("height"),
            stream.get("sample_rate"),
            stream.get("channels"),
        )
        for stream in probe.get("streams", [])
    ], key=lambda item: tuple("" if value is None else str(value) for value in item))


def assert_metadata_sanitized(probe: dict[str, Any], path: Path) -> None:
    def walk(value: Any) -> None:
        if isinstance(value, dict):
            for key, item in value.items():
                normalized = re.sub(r"[^a-z0-9]", "", str(key).lower())
                if normalized in SENSITIVE_METADATA_KEYS:
                    raise ValueError(f"sensitive metadata key {key!r} remains in {path}")
                walk(item)
        elif isinstance(value, list):
            for item in value:
                walk(item)

    walk(probe)


def sanitize_video(source: Path, destination: Path, *, ffmpeg: str, ffprobe: str) -> None:
    temporary = destination.with_name(f".{destination.stem}.tmp{destination.suffix}")
    temporary.unlink(missing_ok=True)
    subprocess.run(
        [
            ffmpeg,
            "-nostdin",
            "-hide_banner",
            "-loglevel",
            "error",
            "-y",
            "-i",
            str(source),
            "-map",
            "0:v?",
            "-map",
            "0:a?",
            "-map",
            "0:s?",
            "-map_metadata",
            "-1",
            "-map_metadata:s",
            "-1",
            "-map_chapters",
            "-1",
            "-metadata",
            "encoder=",
            "-metadata:s",
            "handler_name=",
            "-metadata:s",
            "vendor_id=",
            "-c",
            "copy",
            "-strict",
            "-2",
            "-movflags",
            "+faststart",
            str(temporary),
        ],
        check=True,
    )
    source_probe = probe_media(source, ffprobe)
    output_probe = probe_media(temporary, ffprobe)
    if media_signature(source_probe) != media_signature(output_probe):
        temporary.unlink(missing_ok=True)
        raise ValueError(f"stream signature changed while sanitizing {source}")
    source_duration = float(source_probe.get("format", {}).get("duration") or 0)
    output_duration = float(output_probe.get("format", {}).get("duration") or 0)
    if abs(source_duration - output_duration) > 0.5:
        temporary.unlink(missing_ok=True)
        raise ValueError(f"duration changed while sanitizing {source}")
    assert_metadata_sanitized(output_probe, temporary)
    temporary.replace(destination)


def load_provenance(path: Path | None, expected_videos: set[str]) -> list[dict[str, Any]]:
    if path is None:
        return []
    rows = read_jsonl(path)
    by_video = {str(row.get("video") or "").strip(): row for row in rows}
    if set(by_video) != expected_videos:
        raise ValueError("provenance must contain exactly one row for every released video")
    for video, row in by_video.items():
        if not str(row.get("source") or "").strip() or not str(row.get("license") or "").strip():
            raise ValueError(f"provenance for {video!r} requires source and license")
    return [by_video[name] for name in sorted(by_video)]


def write_dataset_card(output: Path, dataset_id: str, data_license: str, finalized: bool) -> None:
    frontmatter = ""
    if data_license:
        frontmatter = f"---\npretty_name: SVI-Bench\nlicense: {data_license}\nlanguage:\n- zh\n- en\n---\n\n"
    status = "" if finalized else (
        "This is a private staging package. Do not make it public until "
        "provenance and license review is complete."
    )
    text = f"""{frontmatter}# SVI-Bench dataset

SVI-Bench contains 75 streaming-video interaction tasks backed by 71 unique
source videos. Four videos are intentionally reused across task categories;
`media_index.jsonl` is the authoritative task-to-media mapping.

{status}

Use of this dataset is governed by the [SVI-Bench Dataset Terms](DATA_TERMS.md).
It is intended solely for academic research, and commercial use in any form is
prohibited. Copyright in the source videos remains with the respective rights
holders.

## Files

- `SVIBench-开源表.xlsx`: task definitions and dimension-specific grading anchors.
- `interaction-75题/`: 71 metadata-sanitized source videos. Streams are remuxed without
  re-encoding.
- `media_index.jsonl`: mapping from the 75 task IDs to the 71 source videos.
- `annotations/SVI-Pilot-human-ratings.xlsx`: expert pilot labels used for Judge
  alignment.

Use the evaluation code at https://github.com/YidanHAI/SVI-Bench. Pin both the
code commit and the Hugging Face dataset revision when reporting results.

Questions and removal requests may be filed at
https://github.com/YidanHAI/SVI-Bench/issues.

Dataset repository: https://huggingface.co/datasets/{dataset_id}
"""
    (output / "README.md").write_text(text, encoding="utf-8")


def main() -> None:
    args = parse_args()
    workbook = args.workbook.resolve()
    video_dir = args.video_dir.resolve()
    media_index_path = args.media_index.resolve()
    pilot_labels = args.pilot_labels.resolve() if args.pilot_labels else None
    output = args.output.resolve()
    if not workbook.is_file():
        raise FileNotFoundError(workbook)
    if not video_dir.is_dir():
        raise NotADirectoryError(video_dir)
    if pilot_labels is not None and not pilot_labels.is_file():
        raise FileNotFoundError(pilot_labels)
    protected_inputs = tuple(
        item for item in (workbook, video_dir, media_index_path, pilot_labels) if item is not None
    )
    if any(output == item or output in item.parents or item in output.parents for item in protected_inputs):
        raise ValueError("--output must be separate from all input files and directories")

    mapping, sources = load_media_index(media_index_path, video_dir)
    workbook_ids = workbook_task_ids(workbook)
    mapping_ids = [row["id"] for row in mapping]
    if workbook_ids != mapping_ids:
        raise ValueError("workbook task order does not match media_index.jsonl")
    provenance = load_provenance(args.provenance, set(sources))
    if args.finalize and (not provenance or not args.data_license.strip() or pilot_labels is None):
        raise ValueError(
            "--finalize requires complete --provenance, --data-license, and --pilot-labels"
        )

    if args.dry_run:
        print(json.dumps({
            "status": "valid",
            "tasks": len(mapping),
            "unique_videos": len(sources),
            "release_ready": bool(args.finalize),
            "pilot_labels": pilot_labels is not None,
        }, indent=2))
        return

    if output.exists():
        if not args.overwrite:
            raise FileExistsError(f"output exists; pass --overwrite to replace it: {output}")
        shutil.rmtree(output)
    video_output = output / "interaction-75题"
    video_output.mkdir(parents=True)

    sanitized_workbook = output / "SVIBench-开源表.xlsx"
    sanitize_workbook(workbook, sanitized_workbook)
    workbook_task_ids(sanitized_workbook)
    data_terms_source = ROOT / "DATA_TERMS.md"
    if not data_terms_source.is_file():
        raise FileNotFoundError(data_terms_source)
    shutil.copy2(data_terms_source, output / "DATA_TERMS.md")
    if pilot_labels is not None:
        annotation_dir = output / "annotations"
        annotation_dir.mkdir()
        sanitized_pilot = annotation_dir / "SVI-Pilot-human-ratings.xlsx"
        sanitize_workbook(pilot_labels, sanitized_pilot)

    for index, (name, source) in enumerate(sorted(sources.items()), 1):
        print(f"[{index}/{len(sources)}] {name}", flush=True)
        sanitize_video(
            source,
            video_output / name,
            ffmpeg=args.ffmpeg,
            ffprobe=args.ffprobe,
        )

    released_mapping = []
    for row in mapping:
        released_mapping.append({
            "id": row["id"],
            "video": f"interaction-75题/{row['video']}",
        })
    released_index = output / "media_index.jsonl"
    released_index.write_text(
        "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in released_mapping),
        encoding="utf-8",
    )
    if provenance:
        (output / "provenance.jsonl").write_text(
            "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in provenance),
            encoding="utf-8",
        )

    write_dataset_card(output, args.dataset_id, args.data_license.strip(), args.finalize)
    if not args.finalize:
        (output / "PRIVATE_STAGING_ONLY.txt").write_text(
            "Per-video provenance review is incomplete. Keep this dataset private.\n",
            encoding="utf-8",
        )

    print(json.dumps({
        "status": "prepared",
        "output": str(output),
        "tasks": len(mapping),
        "unique_videos": len(sources),
        "release_ready": bool(args.finalize),
    }, indent=2))


if __name__ == "__main__":
    main()
