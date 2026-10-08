import importlib.util
import json
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.packaging.custom import StringProperty


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "prepare_hf_dataset.py"
SPEC = importlib.util.spec_from_file_location("prepare_hf_dataset", SCRIPT)
assert SPEC is not None and SPEC.loader is not None
release = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = release
SPEC.loader.exec_module(release)


def test_released_media_index_has_75_tasks_and_71_unique_videos():
    rows = release.read_jsonl(ROOT / "dataset" / "media_index.jsonl")
    assert len(rows) == 75
    assert len({row["id"] for row in rows}) == 75
    assert len({row["video"] for row in rows}) == 71


def test_media_index_rejects_path_traversal(tmp_path):
    video_dir = tmp_path / "videos"
    video_dir.mkdir()
    outside = tmp_path / "outside.mp4"
    outside.write_bytes(b"video")
    index = tmp_path / "index.jsonl"
    rows = [
        {"id": f"T{number:03d}", "video": f"V{number % 71:03d}.mp4"}
        for number in range(75)
    ]
    for number in range(71):
        (video_dir / f"V{number:03d}.mp4").write_bytes(b"video")
    rows[0]["video"] = "../outside.mp4"
    index.write_text(
        "".join(json.dumps(row) + "\n" for row in rows),
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="must be a basename"):
        release.load_media_index(index, video_dir)


def test_sensitive_video_metadata_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="sensitive metadata"):
        release.assert_metadata_sanitized(
            {"format": {"tags": {"LvMetaInfo": '{"uid":"example"}'}}},
            tmp_path / "video.mp4",
        )


def test_workbook_sanitizer_removes_author_metadata_and_comments(tmp_path):
    source = tmp_path / "source.xlsx"
    destination = tmp_path / "sanitized.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "id"
    sheet["A1"].comment = Comment("internal note", "Private Author")
    workbook.properties.creator = "Private Author"
    workbook.properties.lastModifiedBy = "Private Author"
    workbook.custom_doc_props.append(StringProperty(name="ICV", value="fingerprint"))
    workbook.save(source)
    workbook.close()

    release.sanitize_workbook(source, destination)

    sanitized = load_workbook(destination, keep_links=False)
    try:
        assert sanitized["Sheet"]["A1"].comment is None
        assert sanitized.properties.creator == "SVI-Bench"
        assert sanitized.properties.lastModifiedBy == "SVI-Bench"
        assert list(sanitized.custom_doc_props) == []
    finally:
        sanitized.close()


def test_dataset_card_declares_custom_research_terms(tmp_path):
    release.write_dataset_card(tmp_path, "Danmel02/SVI-bench", "other", False)
    card = (tmp_path / "README.md").read_text(encoding="utf-8")
    assert "license: other" in card
    assert "SVI-Bench Dataset Terms" in card
    assert "solely for academic research" in card
    assert "commercial use in any form is" in card
    assert "SHA256SUMS" not in card
    assert "release_manifest.json" not in card
    assert "per-file hashes" not in card
    assert (ROOT / "DATA_TERMS.md").is_file()
