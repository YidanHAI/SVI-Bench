import json
import re
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import judge_pipeline  # noqa: E402
import judge_workbook_contract as contract  # noqa: E402


TASK_HEADERS = [
    "id",
    "分类",
    "场景标签",
    "入选状态",
    "用户 query",
    "触发点(文字描述)",
    "期望响应窗口",
    "期望关键信息",
]
for dimension in contract.DIMENSION_IDS:
    TASK_HEADERS.extend([
        f"{dimension} 适用",
        f"{dimension} 权重",
        f"{dimension} G/S/B 阈值",
    ])


def inject_wps_data_validation_id(path):
    rewritten = path.with_suffix(".rewritten.xlsx")
    injected = 0
    with zipfile.ZipFile(path, "r") as source:
        with zipfile.ZipFile(rewritten, "w") as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("xl/worksheets/"):
                    payload, count = re.subn(
                        rb"(<dataValidation\b)",
                        rb'\1 id="{wps-non-standard-id}"',
                        payload,
                        count=1,
                    )
                    injected += count
                target.writestr(info, payload)
    assert injected == 1
    rewritten.replace(path)


def make_workbook(path, task_count=1, rubric_note=False, blocking=False):
    workbook = Workbook()
    tasks = workbook.active
    tasks.title = "V1_题目池"
    tasks.append(TASK_HEADERS)
    for index in range(task_count):
        values = {
            "id": f"A{index + 1:04d}",
            "分类": "合成测试",
            "场景标签": "合成场景",
            "入选状态": "入选",
            "用户 query": "看到目标时提醒我",
            "触发点(文字描述)": "目标出现",
            "期望响应窗口": "一秒内",
            "期望关键信息": "目标出现",
        }
        for dimension in contract.DIMENSION_IDS:
            values[f"{dimension} 适用"] = "否" if blocking and dimension == "D3" else "是"
            values[f"{dimension} 权重"] = 0.2
            values[f"{dimension} G/S/B 阈值"] = "G：正确\nS：轻微问题\nB：错误"
        tasks.append([values.get(header) for header in TASK_HEADERS])

    validation = DataValidation(type="list", formula1='"入选,剔除"')
    tasks.add_data_validation(validation)
    validation.add("D2:D100")
    rubric = workbook.create_sheet("评测维度")
    rubric.append(["维度", "定义", "触发条件", "与产品能力的对应"])
    rubric.append(["D1. 主动触发敏感度", "正确触发响应", "目标事件", "自主交互"])
    rubric.append(["D2. 静默正确性", "不该响应时保持安静", "无事件", "自主交互"])
    rubric.append([
        "D3. 响应实时性",
        "触发点到首 token 的延迟；只对有效响应计时，无有效响应记 0",
        "需要响应",
        "实时响应",
    ])
    rubric.append(["D4.响应内容正确性", "响应内容正确", "所有样本", "通用"])
    rubric.append(["D5. 后台委托与记忆", "委托和记忆正确", "复杂请求", "后台委托"])
    if rubric_note:
        rubric["E3"] = "min（）"
    workbook.save(path)
    inject_wps_data_validation_id(path)
    return path


def test_wps_workbook_is_read_without_modifying_source(tmp_path):
    path = make_workbook(tmp_path / "wps.xlsx")
    before = contract.sha256(path)
    workbook = contract.load_workbook_compatible(path, read_only=True, data_only=True)
    assert workbook._joyvl_wps_data_validation_ids_removed == 1
    workbook.close()
    assert list(contract.load_task_specs(path)) == ["A0001"]
    assert contract.sha256(path) == before


def test_headerless_rubric_note_is_non_blocking(tmp_path):
    path = make_workbook(tmp_path / "warning.xlsx", rubric_note=True)
    bundle = contract.build_workbook_v2_contract(path, contract.load_task_specs(path))
    assert bundle["rubric_validation"]["formal_ready"] is True
    assert bundle["rubric_validation"]["issues"][0]["cell"] == "E3"
    assert bundle["rubric_validation"]["issues"][0]["severity"] == "warning"


def test_nonapplicable_dimension_rule_blocks_formal_manifest(tmp_path):
    path = make_workbook(tmp_path / "blocked.xlsx", blocking=True)
    bundle = contract.build_workbook_v2_contract(path, contract.load_task_specs(path))
    errors = [
        item for item in bundle["rubric_validation"]["issues"]
        if item["severity"] == "error"
    ]
    assert bundle["rubric_validation"]["formal_ready"] is False
    assert errors[0]["code"] == "task_nonapplicable_dimension_has_rule"


def test_release_pipeline_accepts_only_ready_75_task_manifest(tmp_path):
    workbook_path = make_workbook(tmp_path / "tasks.xlsx", task_count=75)
    specs = contract.load_task_specs(workbook_path)
    bundle = contract.build_workbook_v2_contract(workbook_path, specs)
    manifest = {
        "schema_version": 2,
        "contains_human_labels": False,
        "source_workbook": str(workbook_path),
        "source_workbook_sha256": contract.sha256(workbook_path),
        "task_count": 75,
        "tasks": [
            {
                "task_spec": contract.bind_contract_to_spec(spec, bundle),
                "evidence_video": {"path": f"/synthetic/{task_id}.mp4"},
            }
            for task_id, spec in specs.items()
        ],
        **bundle,
    }
    path = tmp_path / "manifest.json"
    path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    assert judge_pipeline.validate_manifest(path)["task_count"] == 75

    manifest["contains_human_labels"] = True
    path.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(RuntimeError, match="label-free"):
        judge_pipeline.validate_manifest(path)
